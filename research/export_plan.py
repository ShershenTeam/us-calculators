"""Експортує research/data/page_inventory.csv у Excel-план Page_Plan_12_months.xlsx (корінь проєкту)."""
import csv, os
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

HERE = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.dirname(HERE)
SRC = os.path.join(HERE, "data", "page_inventory.csv")
OUT = os.path.join(ROOT, "Page_Plan_12_months.xlsx")

F = "Arial"
H_FILL = PatternFill("solid", fgColor="1F3864")
H_FONT = Font(name=F, bold=True, color="FFFFFF")
SUB_FILL = PatternFill("solid", fgColor="D9E1F2")
TOT_FILL = PatternFill("solid", fgColor="FFF2CC")
BODY = Font(name=F, size=10)
BOLD = Font(name=F, size=10, bold=True)
THIN = Border(bottom=Side(style="thin", color="BFBFBF"))

TYPES = [("calculator", "Калькулятор"), ("converter", "Конвертер одиниць"), ("value", "Сторінка значення"),
         ("date", "Дата: N днів від сьогодні / тому"), ("countdown", "Відлік / дата свята"),
         ("qa", "Швидка відповідь"), ("state", "Зарплата по штату"), ("generator", "Генератор"),
         ("es", "Іспанська сторінка"), ("hub", "Хаб категорії"), ("service", "Службова сторінка")]
CATS = [("conversions", "Conversions"), ("time-date", "Time & Date"), ("math", "Math"), ("work-pay", "Work & Pay"),
        ("finance", "Finance"), ("construction", "Home & Construction"), ("health", "Health & Fitness"),
        ("business", "Business"), ("everyday", "Everyday"), ("generators", "Generators"), ("science", "Science"),
        ("es", "Español"), ("site", "Службові")]
MONTHS = list(range(13))


def header(ws, row, values, widths=None):
    for i, v in enumerate(values, 1):
        c = ws.cell(row=row, column=i, value=v)
        c.font, c.fill = H_FONT, H_FILL
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    if widths:
        for i, w in enumerate(widths, 1):
            ws.column_dimensions[get_column_letter(i)].width = w


def main():
    rows = list(csv.DictReader(open(SRC, encoding="utf-8")))
    wb = Workbook()

    # ---------- Усі сторінки ----------
    ws = wb.active
    ws.title = "All pages"
    cols = ["#", "Month", "Type", "Category", "URL", "Page / intent", "Competitor traffic/mo", "Min ref. domains",
            "Competitors", "Score", "Best competitor page", "Status", "Notes"]
    header(ws, 1, cols, [6, 9, 12, 14, 46, 38, 14, 11, 34, 8, 60, 12, 30])
    for i, r in enumerate(rows, 2):
        month = int(r["month"]) if r["month"] != "" else "Backlog"
        vals = [i - 1, month, r["type"], r["category"], r["slug"], r["label"],
                int(r["competitor_traffic"] or 0), int(r["min_refdomains"]) if r["min_refdomains"] else None,
                r["competitors"], float(r["score"]), r["example_url"], "Planned" if month != "Backlog" else "Backlog", ""]
        for j, v in enumerate(vals, 1):
            c = ws.cell(row=i, column=j, value=v)
            c.font = BODY
        ws.cell(row=i, column=7).number_format = "#,##0"
        ws.cell(row=i, column=10).number_format = "0.00"
    last = len(rows) + 1
    ws.freeze_panes = "C2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(cols))}{last}"
    rng = lambda col: f"'All pages'!${col}$2:${col}${last}"  # noqa: E731

    # ---------- Зведення по місяцях × тип ----------
    s = wb.create_sheet("Summary", 0)
    s["A1"] = "План сторінок на 12 місяців — зведення"
    s["A1"].font = Font(name=F, size=14, bold=True)
    s["A2"] = "Місяць 0 = запуск (~8-й тиждень робіт). Кількості рахуються формулами з аркуша «All pages»."
    s["A2"].font = Font(name=F, size=9, italic=True, color="595959")
    hdr = ["Тип сторінки"] + [("Запуск" if m == 0 else f"М{m}") for m in MONTHS] + ["Разом за рік", "Беклог"]
    header(s, 4, hdr, [30] + [8] * 13 + [12, 10])
    for i, (code, name) in enumerate(TYPES, 5):
        s.cell(row=i, column=1, value=name).font = BODY
        for j, m in enumerate(MONTHS, 2):
            s.cell(row=i, column=j, value=f'=COUNTIFS({rng("C")},"{code}",{rng("B")},{m})').font = BODY
        s.cell(row=i, column=15, value=f"=SUM(B{i}:N{i})").font = BOLD
        s.cell(row=i, column=16, value=f'=COUNTIFS({rng("C")},"{code}",{rng("B")},"Backlog")').font = BODY
    t = 5 + len(TYPES)
    s.cell(row=t, column=1, value="Нових сторінок за місяць").font = BOLD
    for j in range(2, 17):
        col = get_column_letter(j)
        c = s.cell(row=t, column=j, value=f"=SUM({col}5:{col}{t - 1})")
        c.font, c.fill = BOLD, TOT_FILL
    s.cell(row=t, column=1).fill = TOT_FILL
    s.cell(row=t + 1, column=1, value="Усього сторінок на сайті (наростаючим)").font = BOLD
    for j in range(2, 15):
        col, prev = get_column_letter(j), get_column_letter(j - 1)
        s.cell(row=t + 1, column=j, value=f"={col}{t}" if j == 2 else f"={prev}{t + 1}+{col}{t}").font = BOLD
        s.cell(row=t + 1, column=j).fill = SUB_FILL
    s.cell(row=t + 1, column=1).fill = SUB_FILL
    for rr in range(5, t + 2):
        for cc in range(1, 17):
            s.cell(row=rr, column=cc).border = THIN

    # категорії
    c0 = t + 4
    s.cell(row=c0 - 1, column=1, value="По категоріях").font = Font(name=F, size=12, bold=True)
    header(s, c0, ["Категорія"] + hdr[1:])
    for i, (code, name) in enumerate(CATS, c0 + 1):
        s.cell(row=i, column=1, value=name).font = BODY
        for j, m in enumerate(MONTHS, 2):
            s.cell(row=i, column=j, value=f'=COUNTIFS({rng("D")},"{code}",{rng("B")},{m})').font = BODY
        s.cell(row=i, column=15, value=f"=SUM(B{i}:N{i})").font = BOLD
        s.cell(row=i, column=16, value=f'=COUNTIFS({rng("D")},"{code}",{rng("B")},"Backlog")').font = BODY
        for cc in range(1, 17):
            s.cell(row=i, column=cc).border = THIN
    ct = c0 + 1 + len(CATS)
    s.cell(row=ct, column=1, value="Разом").font = BOLD
    for j in range(2, 17):
        col = get_column_letter(j)
        c = s.cell(row=ct, column=j, value=f"=SUM({col}{c0 + 1}:{col}{ct - 1})")
        c.font, c.fill = BOLD, TOT_FILL
    s.cell(row=ct, column=1).fill = TOT_FILL
    s.cell(row=ct + 1, column=1, value="Перевірка: разом по типах = разом по категоріях").font = Font(name=F, size=9, italic=True)
    s.cell(row=ct + 1, column=15, value=f'=IF(O{ct}=O{t},"OK","ПОМИЛКА")').font = BOLD
    s.freeze_panes = "B5"

    # ---------- Запуск ----------
    L = wb.create_sheet("Launch (month 0)", 1)
    header(L, 1, ["#", "Type", "Category", "URL", "Page / intent", "Competitor traffic/mo", "Min ref. domains"],
           [5, 12, 16, 46, 40, 16, 12])
    launch = [r for r in rows if r["month"] == "0"]
    for i, r in enumerate(launch, 2):
        vals = [i - 1, r["type"], r["category"], r["slug"], r["label"], int(r["competitor_traffic"] or 0),
                int(r["min_refdomains"]) if r["min_refdomains"] else None]
        for j, v in enumerate(vals, 1):
            L.cell(row=i, column=j, value=v).font = BODY
        L.cell(row=i, column=6).number_format = "#,##0"
    L.freeze_panes = "A2"
    L.auto_filter.ref = f"A1:G{len(launch) + 1}"

    # ---------- Легенда ----------
    G = wb.create_sheet("Legend")
    G.column_dimensions["A"].width = 28
    G.column_dimensions["B"].width = 110
    notes = [
        ("Що це", "Перелік усіх сторінок сайту на 12 місяців + беклог. Один рядок = одна сторінка (один намір пошуку)."),
        ("Дата даних", "28.09.2026. Джерело: Ubersuggest, США (топ-сторінки 10 конкурентів: calculator.net, omnicalculator, "
                       "calculatorsoup, inchcalculator, rapidtables, gigacalculator, thecalculatorsite, timeanddate, "
                       "unitconverters.net, widgetly.co; обсяги запитів paycheck по штатах; наші перевірені запити)."),
        ("Month", "0 = запуск (~8-й тиждень робіт), 1–12 = місяці після запуску, Backlog = після 12 місяців."),
        ("Competitor traffic/mo", "Оцінка Ubersuggest: скільки відвідувань на місяць отримує найкраща сторінка конкурента з цим наміром. "
                                  "Для штатних та іспанських сторінок — обсяг пошуку запиту."),
        ("Min ref. domains", "Найменша кількість сайтів-донорів серед сторінок конкурентів з цим наміром. Менше = потрібно менше посилань."),
        ("Score", "Пріоритет: попит (log трафіку) × легкість (тип, донори) × цінність (CPC категорії); YMYL −15%."),
        ("URL", "Чернетка адреси. Остаточний слаг підтверджується перед публікацією за головним запитом (Ubersuggest)."),
        ("Status", "Для відстеження: Planned → In progress → Published. Змінюйте в аркуші All pages."),
        ("Ворота (коли тип стартує)", "Business і генератори — з М2; зарплата по штатах і Health — з М3 (потрібні податковий движок і рецензент); "
                                      "Finance — з М4 (автор-фахівець); Science та іспанська — з М5."),
        ("Пропускна здатність", "Запуск: 130 сторінок + хаби/службові. Далі 60 → 65 → 70 → по 75 на місяць. Калькуляторів не більше 26/міс, "
                                "конвертерів і сторінок значень не більше 32/міс, щоб сайт не став «лише конвертером»."),
        ("Гарантовані місця", "Щомісяця: Math 7, Finance 5 (з М4), Health 3 (з М3), Construction 3, Business 2, генератори 2, "
                              "Work & Pay 2, Everyday 2, Science 2 (з М5), іспанська 3 (з М5), дати/відліки 6."),
        ("Правило сторінок значень", "Сторінка значення («5'7\" in cm», «80 kg to lbs», «90 days from today») — лише якщо в конкурента "
                                     "вона має ≥1 500 відвідувань/міс або запит ≥1 000 пошуків/міс, і на сторінці є власна цінність."),
        ("Як оновлювати", "python research/harvest.py <session.jsonl> → python research/build_inventory.py → python research/export_plan.py"),
    ]
    header(G, 1, ["Поле / правило", "Пояснення"])
    for i, (a, b) in enumerate(notes, 2):
        G.cell(row=i, column=1, value=a).font = BOLD
        c = G.cell(row=i, column=2, value=b)
        c.font, c.alignment = BODY, Alignment(wrap_text=True, vertical="top")
    G.cell(row=len(notes) + 3, column=1, value="Типи сторінок").font = Font(name=F, size=12, bold=True)
    for i, (code, name) in enumerate(TYPES, len(notes) + 4):
        G.cell(row=i, column=1, value=code).font = BOLD
        G.cell(row=i, column=2, value=name).font = BODY

    from openpyxl.workbook.properties import CalcProperties
    wb.calculation = CalcProperties(fullCalcOnLoad=True)   # Excel перераховує формули при відкритті
    wb.save(OUT)
    print("saved", OUT, "rows:", len(rows), "launch:", len(launch))


if __name__ == "__main__":
    main()
